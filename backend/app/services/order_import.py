"""Order import from .xlsx / .csv: safe parsing, column autodetection, product matching."""

import csv
import io
import math
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from decimal import Decimal
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.quantity import MAX_BLOCKS, MAX_PIECES, QuantityError, compute_quantity
from app.models import ProductBarcode, SdProduct
from app.services.catalog import CatalogContext, orderable_items
from app.services.orders import num

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000  # data rows
MAX_SCAN_ROWS = 10_000  # physical rows read from a sheet (blank-row flood guard)
MAX_COLS = 30
HEADER_SCAN_ROWS = 10
MAX_UNZIPPED = 30 * 1024 * 1024
MAX_ZIP_ENTRIES = 200
MAX_ZIP_RATIO = 100
MAX_CELL_LEN = 200
MAX_NUMBER_LEN = 16


class ImportFileError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ImportRow:
    row: int  # 1-based sheet/file row number
    code: str | None
    name: str | None
    blocks: int
    pieces: Decimal
    error: str | None = None


# keyword groups; order matters (see _classify)
_BLOCK_KEYS = ("блок", "blok", "block")
_CODE_KEYS = ("штрих", "shtrix", "barcode", "bar code", "код", "kod", "code", "артикул", "sku")
_QTY_KEYS = ("шт", "qty", "quantity", "кол", "dona", "soni", "miqdor", "pieces", "piece")
_NAME_KEYS = ("наимен", "назв", "товар", "name", "nomi", "mahsulot", "product")


def _classify(header: str) -> str | None:
    h = re.sub(r"\s+", " ", header.strip().lower())
    if not h:
        return None
    for kind, keys in (
        ("blocks", _BLOCK_KEYS),
        ("code", _CODE_KEYS),
        ("pieces", _QTY_KEYS),
        ("name", _NAME_KEYS),
    ):
        if any(k in h for k in keys):
            return kind
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value == int(value) and abs(value) < 1e15:
            return str(int(value))
        return repr(value)[:MAX_CELL_LEN]
    return str(value).strip()[:MAX_CELL_LEN]


# plain decimal notation only: bounded digits, no exponent/inf/nan (``Decimal("1E999999")``
# would otherwise make ``int()`` build a million-digit number)
_NUMBER_RE = re.compile(r"[+-]?\d{1,9}(?:\.\d{1,6})?")


def _to_decimal(value: Any) -> Decimal:
    text = _cell_text(value).replace(" ", "").replace(",", ".")
    if not text:
        return Decimal(0)
    if len(text) > MAX_NUMBER_LEN or not _NUMBER_RE.fullmatch(text):
        raise ValueError("bad_number")
    return Decimal(text)


def _detect_columns(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """Header row index and {kind: column index}; headerless 2-column files are (code, qty)."""
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        found: dict[str, int] = {}
        for col, cell in enumerate(row):
            kind = _classify(_cell_text(cell))
            if kind and kind not in found:
                found[kind] = col
        if ("code" in found or "name" in found) and ("pieces" in found or "blocks" in found):
            return idx, found
    if rows and len([c for c in rows[0] if _cell_text(c)]) >= 2:
        first = rows[0]
        try:
            _to_decimal(first[1])
        except ValueError:
            raise ImportFileError("columns_not_detected", "Could not detect columns") from None
        if _cell_text(first[0]) and _cell_text(first[1]):
            return -1, {"code": 0, "pieces": 1}
    raise ImportFileError(
        "columns_not_detected", "Could not detect columns (need code/barcode and quantity)"
    )


def _rows_to_import(rows: list[list[Any]]) -> list[ImportRow]:
    header_idx, cols = _detect_columns(rows)
    out: list[ImportRow] = []
    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(_cell_text(c) for c in row):
            continue
        if len(out) >= MAX_ROWS:
            raise ImportFileError("too_many_rows", f"At most {MAX_ROWS} rows are allowed")

        def cell(kind: str, row: list[Any] = row) -> Any:
            i = cols.get(kind)
            return row[i] if i is not None and i < len(row) else None

        code = _cell_text(cell("code")) or None
        name = _cell_text(cell("name")) or None
        error: str | None = None
        blocks = 0
        pieces = Decimal(0)
        try:
            b = _to_decimal(cell("blocks"))
            pieces = _to_decimal(cell("pieces"))
            if b != b.to_integral_value():
                error = "bad_blocks"
            blocks = int(b)  # safe: at most 9 digits (see _NUMBER_RE)
            if blocks < 0 or pieces < 0:
                error = "negative_quantity"
            elif blocks > MAX_BLOCKS or pieces > MAX_PIECES:
                error = "quantity_too_large"
            elif not error and blocks == 0 and pieces == 0:
                error = "no_quantity"
        except ValueError:
            error = "bad_number"
        if code is None and name is None:
            error = error or "no_code"
        out.append(ImportRow(offset, code, name, blocks, pieces, error))
    if not out:
        raise ImportFileError("empty_file", "No data rows found")
    return out


def _check_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            infos = z.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise ImportFileError("bad_file", "Archive has too many entries")
            total = 0
            for info in infos:
                total += info.file_size
                if info.compress_size and info.file_size / info.compress_size > MAX_ZIP_RATIO:
                    raise ImportFileError("bad_file", "Suspicious compression ratio")
            if total > MAX_UNZIPPED:
                raise ImportFileError("bad_file", "Uncompressed content is too large")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("bad_file", "Not a valid .xlsx file") from exc


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook  # heavy import, only needed here

    _check_zip(content)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("bad_file", "Could not read the spreadsheet") from exc
    try:
        ws = wb.worksheets[0]
        rows: list[list[Any]] = []
        non_empty = 0
        for values in ws.iter_rows(values_only=True, max_row=MAX_SCAN_ROWS, max_col=MAX_COLS):
            row = list(values)
            rows.append(row)
            if any(_cell_text(c) for c in row):
                non_empty += 1
                if non_empty > MAX_ROWS + HEADER_SCAN_ROWS:
                    raise ImportFileError("too_many_rows", f"At most {MAX_ROWS} rows are allowed")
        return rows
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("bad_file", "Could not read the spreadsheet") from exc
    finally:
        wb.close()


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("bad_file", "Unsupported text encoding")


def _read_csv(content: bytes) -> list[list[Any]]:
    text = _decode(content)
    if "\x00" in text:
        raise ImportFileError("bad_file", "Not a text file")
    first = text.split("\n", 1)[0]
    delimiter = max((",", ";", "\t"), key=first.count)
    rows: list[list[Any]] = []
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
            if i >= MAX_SCAN_ROWS:
                raise ImportFileError("too_many_rows", f"At most {MAX_ROWS} rows are allowed")
            rows.append([c[:MAX_CELL_LEN] for c in row[:MAX_COLS]])
    except csv.Error as exc:
        raise ImportFileError("bad_file", "Could not parse the CSV file") from exc
    return rows


def parse_upload(filename: str, content: bytes) -> list[ImportRow]:
    """Parse an uploaded .xlsx/.csv into rows. Raises ``ImportFileError`` for bad input."""
    if len(content) > MAX_BYTES:
        raise ImportFileError("file_too_large", "File is larger than 2 MB")
    if not content:
        raise ImportFileError("empty_file", "File is empty")
    name = (filename or "").lower()
    is_zip = content[:4] == b"PK\x03\x04"
    if name.endswith(".xlsx"):
        if not is_zip:
            raise ImportFileError("bad_file", "Not a valid .xlsx file")
        rows = _read_xlsx(content)
    elif name.endswith(".csv"):
        if is_zip:
            raise ImportFileError("bad_file", "Not a CSV file")
        rows = _read_csv(content)
    else:
        raise ImportFileError("unsupported_type", "Only .xlsx and .csv files are supported")
    return _rows_to_import(rows)


# ------------------------------------------------------------------------------------ matching


def _lower(values: Iterable[str]) -> list[str]:
    return list({v.lower() for v in values})


async def match_rows(db: AsyncSession, cx: CatalogContext, rows: list[ImportRow]) -> dict[str, Any]:
    """Preview: matched (unique orderable product), ambiguous (several), unmatched (+reason)."""
    codes = [r.code for r in rows if r.code and not r.error]
    names = [r.name for r in rows if r.name and not r.code and not r.error]
    by_barcode: dict[str, set[int]] = {}
    by_code: dict[str, set[int]] = {}
    by_sd: dict[str, set[int]] = {}
    by_name: dict[str, set[int]] = {}
    if codes:
        for code, pid in (
            await db.execute(
                select(ProductBarcode.barcode, ProductBarcode.product_id).where(
                    ProductBarcode.barcode.in_(codes)
                )
            )
        ).all():
            by_barcode.setdefault(code, set()).add(pid)
        for c1c, pid in (
            await db.execute(
                select(SdProduct.code_1c, SdProduct.id).where(SdProduct.code_1c.in_(codes))
            )
        ).all():
            if c1c:
                by_code.setdefault(c1c, set()).add(pid)
        for sd_id, pid in (
            await db.execute(
                select(SdProduct.sd_id, SdProduct.id).where(SdProduct.sd_id.in_(codes))
            )
        ).all():
            by_sd.setdefault(sd_id, set()).add(pid)
    if names:
        for lname, pid in (
            await db.execute(
                select(func.lower(SdProduct.name), SdProduct.id).where(
                    func.lower(SdProduct.name).in_(_lower(names))
                )
            )
        ).all():
            by_name.setdefault(lname, set()).add(pid)

    candidates: dict[int, set[int]] = {}
    for r in rows:
        if r.error:
            continue
        if r.code:
            ids = by_barcode.get(r.code) or by_code.get(r.code) or by_sd.get(r.code) or set()
        else:
            ids = by_name.get((r.name or "").lower(), set())
        candidates[r.row] = ids
    all_ids = sorted({i for ids in candidates.values() for i in ids})
    items = await orderable_items(db, cx, all_ids)
    packs = (
        dict(
            (
                await db.execute(
                    select(SdProduct.id, SdProduct.pack_quantity).where(SdProduct.id.in_(all_ids))
                )
            ).all()
        )
        if all_ids
        else {}
    )

    matched: dict[int, dict[str, Any]] = {}
    ambiguous: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []

    def miss(r: ImportRow, reason: str) -> None:
        unmatched.append({"row": r.row, "code": r.code, "name": r.name, "reason": reason})

    for r in rows:
        if r.error:
            miss(r, r.error)
            continue
        ids = candidates[r.row]
        if not ids:
            miss(r, "not_found")
            continue
        ok_ids = sorted(i for i in ids if i in items)
        if not ok_ids:
            miss(r, "not_orderable")
            continue
        if len(ok_ids) > 1:
            ambiguous.append(
                {
                    "row": r.row,
                    "code": r.code,
                    "name": r.name,
                    "blocks": r.blocks,
                    "pieces": num(r.pieces),
                    "candidates": [items[i] for i in ok_ids],
                }
            )
            continue
        pid = ok_ids[0]
        try:
            compute_quantity(r.blocks, r.pieces, packs.get(pid))
        except QuantityError as exc:
            miss(r, exc.code)
            continue
        entry = matched.get(pid)
        if entry is None:
            matched[pid] = {**items[pid], "rows": [r.row], "blocks": r.blocks, "pieces": r.pieces}
        else:  # the same product on several rows: quantities add up
            entry["rows"].append(r.row)
            entry["blocks"] += r.blocks
            entry["pieces"] += r.pieces
    out_matched: list[dict[str, Any]] = []
    for pid, entry in matched.items():
        quantity = compute_quantity(entry["blocks"], entry["pieces"], packs.get(pid))
        out_matched.append({**entry, "pieces": num(entry["pieces"]), "quantity": num(quantity)})
    return {
        "rows": len(rows),
        "matched": out_matched,
        "ambiguous": ambiguous,
        "unmatched": unmatched,
    }
