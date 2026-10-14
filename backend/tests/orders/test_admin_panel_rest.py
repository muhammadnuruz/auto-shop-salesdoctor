"""Admin panel remainder: sales reports, link statistics, runtime settings, Telegram notifier,
DB-driven sync dispatcher, audit log browser."""

import io
import json
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from typing import Any

import httpx
import pytest
import respx
from sqlalchemy import select

from app.core import clock
from app.models import AppSetting, AuditLog, ShopOrder, SyncRun
from app.services import app_settings, notifier
from app.workers import settings as worker
from tests.api.conftest import Api, ApiFactory
from tests.orders.conftest import World

TOKEN = "123456789:AAH" + "x" * 32
TG = f"https://api.telegram.org/bot{TOKEN}/sendMessage"


async def place(world: World, lines: list[tuple[str, int, int | str]]) -> dict[str, Any]:
    order = await world.draft(lines)
    r = await world.submit(order)
    assert r.status_code == 200, r.text
    body: dict[str, Any] = r.json()
    return body


async def login_as(app_factory: ApiFactory, world: World, username: str, role: str) -> Api:
    await world.seed.admin(username, role=role)
    api = await app_factory()
    assert (await api.admin_login(username)).status_code == 200
    return api


# ------------------------------------------------------------------------------------ reports


async def test_sales_report_groupings_and_totals(world: World) -> None:
    await place(world, [("pb", 0, 2)])  # 2 x 500
    await place(world, [("pa", 0, 3), ("pb", 0, 1)])  # 3 x 1000 + 500
    await world.draft([("pa", 0, 1)])  # draft: never counted
    url = "/api/admin/reports/sales"

    r = (await world.admin.get(url, params={"group_by": "client"})).json()
    assert r["totals"] == {"orders": 2, "qty": 6, "sum": 4500}
    assert r["rows"] == [
        {"key": str(world.client.id), "name": "Тест Дилшод", "orders": 2, "qty": 6, "sum": 4500}
    ]
    assert r["truncated"] is False

    rows = (await world.admin.get(url, params={"group_by": "product"})).json()["rows"]
    assert [(x["key"], x["orders"], x["qty"], x["sum"]) for x in rows] == [
        ("pa", 1, 3, 3000),
        ("pb", 2, 3, 1500),
    ]
    rows = (await world.admin.get(url, params={"group_by": "category"})).json()["rows"]
    assert {(x["key"], x["name"], x["sum"]) for x in rows} == {
        ("c1", "Cat 1", 3000),
        ("c2", "Cat 2", 1500),
    }
    rows = (await world.admin.get(url, params={"group_by": "link"})).json()["rows"]
    assert rows[0]["key"] == str(world.link.id) and rows[0]["orders"] == 2
    rows = (await world.admin.get(url, params={"group_by": "agent"})).json()["rows"]
    assert rows[0]["key"] == "d0_11" and rows[0]["name"] == "AGENT11"
    rows = (await world.admin.get(url, params={"group_by": "day"})).json()["rows"]
    assert len(rows) == 1 and rows[0]["orders"] == 2 and rows[0]["key"] == rows[0]["name"]

    # filters
    r = (await world.admin.get(url, params={"link_id": world.link.id + 999})).json()
    assert r["totals"]["orders"] == 0 and r["rows"] == []
    r = await world.admin.get(url, params={"date_from": "2000-01-01", "date_to": "2000-01-31"})
    assert r.json()["totals"]["orders"] == 0


async def test_sales_report_validation(world: World) -> None:
    url = "/api/admin/reports/sales"
    r = await world.admin.get(url, params={"date_from": "2026-02-01", "date_to": "2026-01-01"})
    assert r.status_code == 422
    r = await world.admin.get(url, params={"date_from": "2024-01-01", "date_to": "2026-01-01"})
    assert r.status_code == 422
    assert (await world.admin.get(url, params={"group_by": "nope"})).status_code == 422


async def test_sales_report_xlsx_is_audited_and_formula_safe(world: World) -> None:
    from openpyxl import load_workbook

    world.client.name = "=HYPERLINK(1)"
    await world.session.commit()
    await place(world, [("pb", 0, 2)])
    r = await world.admin.get("/api/admin/reports/sales.xlsx", params={"group_by": "client"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[1] == ["Key", "Name", "Orders", "Qty", "Sum"]
    assert values[2][1] == "'=HYPERLINK(1)" and values[2][4] == 1000
    assert values[-1][1] == "TOTAL"
    a = (
        await world.session.scalars(select(AuditLog).where(AuditLog.action == "report.export"))
    ).one()
    assert a.entity_id == "sales:client" and (a.after or {})["rows"] == 1


async def test_reports_permissions(world: World, app_factory: ApiFactory) -> None:
    viewer = await login_as(app_factory, world, "viewer1", "viewer")
    assert (await viewer.get("/api/admin/reports/sales")).status_code == 200
    op = await login_as(app_factory, world, "op1", "operator")
    assert (await op.get("/api/admin/reports/sales")).status_code == 403
    assert (await op.get("/api/admin/reports/sales.xlsx")).status_code == 403


# --------------------------------------------------------------------------------- link stats


async def test_link_stats_counts_orders_and_logins(world: World) -> None:
    await place(world, [("pb", 0, 2)])
    await world.draft([("pa", 0, 1)])
    s = (await world.admin.get(f"/api/admin/links/{world.link.id}/stats?days=7")).json()
    assert s["active_sessions"] >= 1 and s["unique_clients"] == 1
    assert s["orders"]["all_time"] == {"orders": 1, "sum": 1000}
    assert s["orders"]["period"]["orders"] == 1 and s["orders"]["period"]["sum"] == 1000
    assert s["orders"]["period"]["logins"] >= 1
    assert s["orders"]["by_state"]["sent"]["orders"] == 1
    assert s["orders"]["by_state"]["draft"]["orders"] == 1
    assert len(s["daily"]) == 7 and s["daily"][-1]["orders"] == 1
    assert s["top_clients"][0]["client_id"] == world.client.id
    assert (
        await world.admin.get(f"/api/admin/links/{world.link.id}/stats?days=0")
    ).status_code == 422


# ----------------------------------------------------------------------------------- settings


async def test_settings_roundtrip_never_exposes_token(world: World) -> None:
    r = await world.admin.get("/api/admin/settings")
    assert r.status_code == 200
    body = r.json()
    assert body["sync_intervals"]["stock"]["override"] is None
    assert body["telegram"]["token_configured"] is False
    assert body["sd"]["host"] == "example.salesdoc.io" and "login" not in body["sd"]

    r = await world.admin.put(
        "/api/admin/settings",
        {
            "sync_intervals": {"stock": 120, "clients": 3600},
            "phone_max_clients": 5,
            "telegram": {"enabled": True, "bot_token": TOKEN, "chat_ids": ["-1001", " @my_chan "]},
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert TOKEN not in r.text
    assert body["sync_intervals"]["stock"] == {
        **body["sync_intervals"]["stock"],
        "override": 120,
        "effective": 120,
    }
    assert body["phone_max_clients"]["effective"] == 5
    assert body["telegram"]["token_configured"] and body["telegram"]["active"]
    assert body["telegram"]["chat_ids"] == ["-1001", "@my_chan"]

    stored = await world.session.get(AppSetting, "telegram", populate_existing=True)
    assert stored is not None and TOKEN not in str(stored.value)
    audits = (
        await world.session.scalars(select(AuditLog).where(AuditLog.action == "settings.update"))
    ).all()
    assert audits and all(TOKEN not in str(a.after) for a in audits)
    assert (audits[-1].after or {})["telegram"]["bot_credential"] == "changed"

    # null removes an override, reset flag removes the phone override, "" clears the token
    r = await world.admin.put(
        "/api/admin/settings",
        {
            "sync_intervals": {"stock": None},
            "reset_phone_max_clients": True,
            "telegram": {"bot_token": ""},
        },
    )
    body = r.json()
    assert body["sync_intervals"]["stock"]["override"] is None
    assert body["sync_intervals"]["clients"]["override"] == 3600
    assert body["phone_max_clients"]["override"] is None
    assert body["telegram"]["token_configured"] is False and body["telegram"]["active"] is False


@pytest.mark.parametrize(
    "bad",
    [
        {"sync_intervals": {"stock": 5}},
        {"sync_intervals": {"nope": 600}},
        {"sync_intervals": {"clients": 100}},
        {"phone_max_clients": 0},
        {"telegram": {"bot_token": "not-a-token"}},
        {"telegram": {"chat_ids": ["abc def"]}},
        {"telegram": {"events": ["everything"]}},
    ],
)
async def test_settings_validation(world: World, bad: dict[str, Any]) -> None:
    assert (await world.admin.put("/api/admin/settings", bad)).status_code == 422


async def test_settings_permission(world: World, app_factory: ApiFactory) -> None:
    viewer = await login_as(app_factory, world, "viewer2", "viewer")
    assert (await viewer.get("/api/admin/settings")).status_code == 403
    assert (await viewer.put("/api/admin/settings", {"phone_max_clients": 4})).status_code == 403
    adm = await login_as(app_factory, world, "adm2", "admin")
    assert (await adm.get("/api/admin/settings")).status_code == 200


async def test_phone_threshold_override_is_used(world: World) -> None:
    sm = world.sync.settings
    assert await app_settings.phone_max_clients(world.session, sm) == sm.PHONE_MAX_CLIENTS
    await world.admin.put("/api/admin/settings", {"phone_max_clients": 7})
    assert await app_settings.phone_max_clients(world.session, sm) == 7


async def test_telegram_test_endpoint(world: World, mock_sd: respx.MockRouter) -> None:
    r = await world.admin.post("/api/admin/settings/telegram/test")
    assert r.status_code == 409
    await world.admin.put(
        "/api/admin/settings",
        {"telegram": {"enabled": True, "bot_token": TOKEN, "chat_ids": ["1", "2"]}},
    )
    route = mock_sd.post(TG).mock(
        side_effect=[httpx.Response(200, json={"ok": True}), httpx.Response(400, json={})]
    )
    world.admin.http._transport.app.state.http = httpx.AsyncClient()  # type: ignore[attr-defined]
    r = await world.admin.post("/api/admin/settings/telegram/test")
    assert r.status_code == 200
    assert r.json() == {"results": {"1": "ok", "2": "error: http 400"}, "ok": False}
    assert route.call_count == 2 and TOKEN not in r.text


# ----------------------------------------------------------------------------------- notifier


async def _configure_telegram(world: World, **extra: Any) -> None:
    r = await world.admin.put(
        "/api/admin/settings",
        {"telegram": {"enabled": True, "bot_token": TOKEN, "chat_ids": ["42"], **extra}},
    )
    assert r.status_code == 200


async def test_notifier_starts_at_now_and_delivers_new_orders_once(
    world: World, mock_sd: respx.MockRouter
) -> None:
    await place(world, [("pb", 0, 1)])  # before the first tick: history, never replayed
    route = mock_sd.post(TG).mock(return_value=httpx.Response(200, json={"ok": True}))
    s = world.sync.settings
    async with httpx.AsyncClient() as http:
        assert await notifier.tick(world.session, http, s) == 0  # initialises the cursor
        await _configure_telegram(world)
        await place(world, [("pb", 0, 2)])
        code = await world.session.scalar(
            select(ShopOrder.code_1c).order_by(ShopOrder.id.desc()).limit(1)
        )
        assert await notifier.tick(world.session, http, s) == 1
        assert await notifier.tick(world.session, http, s) == 0
    assert route.call_count == 1
    sent = json.loads(route.calls[0].request.content)
    assert sent["chat_id"] == "42"
    assert code and code in sent["text"] and "1 000" in sent["text"]


async def test_notifier_retries_transient_failures(
    world: World, mock_sd: respx.MockRouter, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = clock.utcnow()
    async with httpx.AsyncClient() as http:
        await _tick_at(world, http, monkeypatch, now)
        await _configure_telegram(world)
        await place(world, [("pb", 0, 1)])
        route = mock_sd.post(TG).mock(
            side_effect=[httpx.ConnectError("boom"), httpx.Response(200, json={"ok": True})]
        )
        assert await _tick_at(world, http, monkeypatch, now) == 0
        assert await _tick_at(world, http, monkeypatch, now + timedelta(seconds=10)) == 0  # backoff
        assert await _tick_at(world, http, monkeypatch, now + timedelta(seconds=31)) == 1
        assert await _tick_at(world, http, monkeypatch, now + timedelta(seconds=62)) == 0
    assert route.call_count == 2


async def test_notifier_respects_disabled_events_and_advances(
    world: World, mock_sd: respx.MockRouter
) -> None:
    route = mock_sd.post(TG).mock(return_value=httpx.Response(200, json={"ok": True}))
    s = world.sync.settings
    async with httpx.AsyncClient() as http:
        await notifier.tick(world.session, http, s)
        await _configure_telegram(world, events=["sync_error"])
        await place(world, [("pb", 0, 1)])
        assert await notifier.tick(world.session, http, s) == 0
        await world.admin.put("/api/admin/settings", {"telegram": {"events": ["new_order"]}})
        assert await notifier.tick(world.session, http, s) == 0  # already passed, not replayed
    assert route.call_count == 0


async def test_notifier_sync_error_transitions_only(
    world: World, mock_sd: respx.MockRouter
) -> None:
    route = mock_sd.post(TG).mock(return_value=httpx.Response(200, json={"ok": True}))
    s = world.sync.settings
    base = clock.utcnow() - timedelta(minutes=10)
    async with httpx.AsyncClient() as http:
        await notifier.tick(world.session, http, s)
        await _configure_telegram(world)

        def run(i: int, status: str) -> SyncRun:
            at = base + timedelta(seconds=10 * (i + 1))
            return SyncRun(
                entity="stock", started_at=at, finished_at=at, status=status, error=f"err-{i}"
            )

        world.session.add_all(
            [run(0, "error"), run(1, "error"), run(2, "error"), run(3, "ok"), run(4, "ok")]
        )
        running = SyncRun(entity="prices", started_at=base, status="running")
        world.session.add(running)
        await world.session.commit()
        assert await notifier.tick(world.session, http, s) == 2  # failure + recovery
        assert await notifier.tick(world.session, http, s) == 0

        # the long-running run finishes later with an earlier-looking id: still reported
        running.status, running.error = "error", "prices broke"
        running.finished_at = base + timedelta(seconds=100)
        await world.session.commit()
        assert await notifier.tick(world.session, http, s) == 1

        # a run that finished just now is held back by the commit lag, then delivered
        fresh = SyncRun(
            entity="clients",
            started_at=clock.utcnow(),
            finished_at=clock.utcnow(),
            status="error",
            error="clients broke",
        )
        world.session.add(fresh)
        await world.session.commit()
        assert await notifier.tick(world.session, http, s) == 0
    texts = [json.loads(c.request.content)["text"] for c in route.calls]
    assert "«stock»" in texts[0] and "err-0" in texts[0] and "err-1" not in texts[0]
    assert "восстановлена" in texts[1]
    assert "prices broke" in texts[2]


async def _tick_at(
    world: World, http: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch, at: datetime
) -> int:
    monkeypatch.setattr(notifier, "clock", SimpleNamespace(utcnow=lambda: at))
    return await notifier.tick(world.session, http, world.sync.settings)


async def test_notifier_per_chat_progress_and_retry_after(
    world: World, mock_sd: respx.MockRouter, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = clock.utcnow()
    calls: list[str] = []
    limited = {"left": 1}  # the second chat is rate limited once

    def handler(request: httpx.Request) -> httpx.Response:
        chat = json.loads(request.content)["chat_id"]
        calls.append(chat)
        if chat == "@limited" and limited["left"] > 0:
            limited["left"] -= 1
            return httpx.Response(429, json={"ok": False, "parameters": {"retry_after": 120}})
        return httpx.Response(200, json={"ok": True})

    mock_sd.post(TG).mock(side_effect=handler)
    async with httpx.AsyncClient() as http:
        await _tick_at(world, http, monkeypatch, now)
        await world.admin.put(
            "/api/admin/settings",
            {"telegram": {"enabled": True, "bot_token": TOKEN, "chat_ids": ["111", "@limited"]}},
        )
        await place(world, [("pb", 0, 1)])
        assert await _tick_at(world, http, monkeypatch, now) == 0
        assert calls == ["111", "@limited"]
        # inside retry_after: nothing is sent
        assert await _tick_at(world, http, monkeypatch, now + timedelta(seconds=60)) == 0
        assert calls == ["111", "@limited"]
        # afterwards only the chat that missed it gets it: no duplicate for 111
        assert await _tick_at(world, http, monkeypatch, now + timedelta(seconds=130)) == 1
        assert calls == ["111", "@limited", "@limited"]
        h = (await world.admin.get("/api/admin/settings")).json()["telegram"]["health"]
        assert h["last_ok_at"] and "429" in h["last_error"]


async def test_notifier_bad_token_blocks_and_bad_chat_is_skipped(
    world: World, mock_sd: respx.MockRouter, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = clock.utcnow()
    async with httpx.AsyncClient() as http:
        await _tick_at(world, http, monkeypatch, now)
        await _configure_telegram(world)
        await place(world, [("pb", 0, 1)])
        route = mock_sd.post(TG).mock(return_value=httpx.Response(401, json={"ok": False}))
        assert await _tick_at(world, http, monkeypatch, now) == 0
        assert await _tick_at(world, http, monkeypatch, now + timedelta(minutes=1)) == 0
        assert route.call_count == 2  # retried, not dropped
        h = (await world.admin.get("/api/admin/settings")).json()["telegram"]["health"]
        assert "token rejected" in h["last_error"] and TOKEN not in str(h)

        route.mock(return_value=httpx.Response(200, json={"ok": True}))
        assert await _tick_at(world, http, monkeypatch, now + timedelta(minutes=2)) == 1

        # a chat that refuses (400) is skipped for that message instead of blocking forever
        await place(world, [("pb", 0, 1)])
        route.mock(return_value=httpx.Response(400, json={"ok": False}))
        assert await _tick_at(world, http, monkeypatch, now + timedelta(minutes=3)) == 1
        assert await _tick_at(world, http, monkeypatch, now + timedelta(minutes=4)) == 0


async def test_notifier_drops_stale_and_limits_per_tick(
    world: World, mock_sd: respx.MockRouter, monkeypatch: pytest.MonkeyPatch
) -> None:
    route = mock_sd.post(TG).mock(return_value=httpx.Response(200, json={"ok": True}))
    now = clock.utcnow()
    async with httpx.AsyncClient() as http:
        await _tick_at(world, http, monkeypatch, now)
        await _configure_telegram(world)
        await place(world, [("pb", 0, 1)])
        # seen 7 hours later (e.g. after an outage): too old, dropped silently
        assert await _tick_at(world, http, monkeypatch, now + timedelta(hours=7)) == 0
        assert route.call_count == 0

        base = clock.utcnow() - timedelta(minutes=30)
        for i in range(notifier.MAX_PER_TICK + 3):
            at = base + timedelta(seconds=i)
            status = "error" if i % 2 == 0 else "ok"
            world.session.add(SyncRun(entity="stock", started_at=at, finished_at=at, status=status))
        await world.session.commit()
        assert await _tick_at(world, http, monkeypatch, now) == notifier.MAX_PER_TICK
        assert await _tick_at(world, http, monkeypatch, now) == 3


async def test_notifier_picks_up_events_committed_out_of_order(
    world: World, mock_sd: respx.MockRouter
) -> None:
    from app.models import ShopOrderEvent

    route = mock_sd.post(TG).mock(return_value=httpx.Response(200, json={"ok": True}))
    s = world.sync.settings
    async with httpx.AsyncClient() as http:
        await notifier.tick(world.session, http, s)
        await _configure_telegram(world)
        await place(world, [("pb", 0, 1)])
        order_id = await world.session.scalar(select(ShopOrder.id).order_by(ShopOrder.id.desc()))
        top = await world.session.scalar(
            select(ShopOrderEvent.id).order_by(ShopOrderEvent.id.desc())
        )
        assert await notifier.tick(world.session, http, s) == 1
        # an older event (lower id than the cursor) only becomes relevant now — the same
        # situation as a transaction that committed after a newer one was processed
        older = (
            await world.session.scalars(
                select(ShopOrderEvent).where(
                    ShopOrderEvent.order_id == order_id, ShopOrderEvent.event == "submitted"
                )
            )
        ).one()
        assert older.id < (top or 0)
        older.event = "outcome_unknown"
        await world.session.commit()
        assert await notifier.tick(world.session, http, s) == 1
        assert await notifier.tick(world.session, http, s) == 0
    assert route.call_count == 2


async def test_notifier_maps_order_problems() -> None:
    from app.models import ShopOrderEvent

    order = ShopOrder(code_1c="WEB-000001", total_sum=1234567)

    def ev(event: str, payload: dict[str, Any] | None = None) -> ShopOrderEvent:
        return ShopOrderEvent(event=event, payload=payload)

    def kind(e: ShopOrderEvent) -> str | None:
        n = notifier._order_notice(e, order, "Shop", "slug")
        return n.kind if n else None

    assert kind(ev("sd_ok")) == "new_order"
    assert kind(ev("reconciled", {"op": "submit"})) == "new_order"
    assert kind(ev("reconciled", {"op": "cancel"})) is None
    assert kind(ev("outcome_unknown")) == "outcome_unknown"
    assert kind(ev("sd_error", {"message": notifier.EXHAUSTED_MESSAGE})) == "reconcile_exhausted"
    assert kind(ev("sd_error", {"errors": ["x"]})) == "order_failed"
    n = notifier._order_notice(ev("sd_ok"), order, "Shop", "slug")
    assert n is not None and "1 234 567" in n.text


# --------------------------------------------------------------------------------- dispatcher


class FakeArq:
    """ArqRedis stand-in: real (fake)redis for the slot claim, recorded enqueues."""

    def __init__(self, redis: Any) -> None:
        self.redis = redis
        self.calls: list[str] = []

    async def set(self, *a: Any, **kw: Any) -> Any:
        return await self.redis.set(*a, **kw)

    async def enqueue_job(self, name: str, *, _job_id: str) -> object:
        self.calls.append(name)
        return object()


def _at(monkeypatch: pytest.MonkeyPatch, t: float) -> None:
    monkeypatch.setattr(worker, "time", SimpleNamespace(time=lambda: float(t)))


T0 = 1_800_000_000  # divisible by 60, 300, 600, 1800 and 3600 -> every slot starts here


async def test_dispatcher_uses_db_intervals_and_dedupes(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    arq, other_worker = FakeArq(world.redis), FakeArq(world.redis)
    ctx: dict[str, Any] = {"sync": world.sync, "redis": arq}
    _at(monkeypatch, T0)
    assert set(await worker.dispatch_syncs(ctx)) == set(app_settings.SYNC_ENTITIES)
    assert await worker.dispatch_syncs(ctx) == []  # same slot: claimed
    assert await worker.dispatch_syncs({**ctx, "redis": other_worker}) == []  # other worker

    _at(monkeypatch, T0 + 60)
    assert await worker.dispatch_syncs(ctx) == ["stock"]  # default stock interval = 60 s

    # the stock sync queued at T0+60 actually ran
    world.session.add(
        SyncRun(entity="stock", started_at=datetime.fromtimestamp(T0 + 61, UTC), status="ok")
    )
    await world.session.commit()
    r = await world.admin.put("/api/admin/settings", {"sync_intervals": {"stock": 300}})
    assert r.status_code == 200, r.text
    _at(monkeypatch, T0 + 120)
    assert await worker.dispatch_syncs(ctx) == []  # new 300 s slot is covered by that run
    _at(monkeypatch, T0 + 300 + 5)
    assert await worker.dispatch_syncs(ctx) == ["stock", "balances"]


async def test_dispatcher_late_tick_still_runs_the_slot(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    arq = FakeArq(world.redis)
    ctx: dict[str, Any] = {"sync": world.sync, "redis": arq}
    _at(monkeypatch, T0 + 1800 - 1)  # previous slots
    await worker.dispatch_syncs(ctx)
    _at(monkeypatch, T0 + 1800 + 95)  # the tick at the slot start was missed (busy worker)
    assert "clients" in await worker.dispatch_syncs(ctx)


async def test_dispatcher_skips_slot_covered_by_a_run(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    world.session.add(
        SyncRun(
            entity="clients",
            started_at=datetime.fromtimestamp(T0 + 10, UTC),
            status="ok",
        )
    )
    await world.session.commit()
    arq = FakeArq(world.redis)
    _at(monkeypatch, T0 + 20)  # e.g. the startup full sync already ran in this slot
    queued = await worker.dispatch_syncs({"sync": world.sync, "redis": arq})
    assert "clients" not in queued and "products" in queued


def test_worker_registers_dispatch_and_notify() -> None:
    names = {getattr(f, "__name__", "") for f in worker.WorkerSettings.functions}
    assert {"dispatch_syncs", "notify_tick"} <= names
    assert {f"sync_{e}" for e in app_settings.SYNC_ENTITIES} <= names
    cron_names = {c.name for c in worker.WorkerSettings.cron_jobs}
    assert any("notify_tick" in n for n in cron_names)


# -------------------------------------------------------------------------------------- audit


async def test_audit_list_filters_and_facets(world: World, app_factory: ApiFactory) -> None:
    await place(world, [("pb", 0, 1)])
    await world.admin.put("/api/admin/settings", {"phone_max_clients": 4})
    r = (await world.admin.get("/api/admin/audit")).json()
    assert r["total"] >= 2
    actions = {i["action"] for i in r["items"]}
    assert {"settings.update", "order.submit"} <= actions
    ats = [datetime.fromisoformat(i["at"]) for i in r["items"]]
    assert ats == sorted(ats, reverse=True)

    r = (await world.admin.get("/api/admin/audit", params={"action": "settings.update"})).json()
    assert r["total"] == 1 and r["items"][0]["actor_name"] == "root"
    r = (await world.admin.get("/api/admin/audit", params={"action": "order."})).json()
    assert r["total"] >= 1 and all(i["action"].startswith("order.") for i in r["items"])
    r = (await world.admin.get("/api/admin/audit", params={"actor_type": "shop"})).json()
    assert all(i["actor_type"] == "shop" for i in r["items"])
    r = (
        await world.admin.get(
            "/api/admin/audit",
            params={"date_from": "2000-01-01T00:00:00Z", "date_to": "2000-01-02T00:00:00Z"},
        )
    ).json()
    assert r == {"total": 0, "items": []}
    assert (await world.admin.get("/api/admin/audit?date_from=2000-01-01")).status_code == 422

    f = (await world.admin.get("/api/admin/audit/facets")).json()
    assert "settings.update" in f["actions"] and "settings" in f["entity_types"]
    assert any(a["username"] == "root" for a in f["actors"])

    op = await login_as(app_factory, world, "op3", "operator")
    assert (await op.get("/api/admin/audit")).status_code == 403
    assert (await op.get("/api/admin/audit/facets")).status_code == 403


# ------------------------------------------------------------------------ review regressions


async def test_link_logins_survive_session_purge(world: World) -> None:
    from sqlalchemy import delete

    from app.models import ShopSession

    before = (await world.admin.get(f"/api/admin/links/{world.link.id}/stats")).json()
    await world.session.execute(delete(ShopSession).where(ShopSession.link_id == world.link.id))
    await world.session.commit()
    after = (await world.admin.get(f"/api/admin/links/{world.link.id}/stats")).json()
    assert after["orders"]["period"]["logins"] == before["orders"]["period"]["logins"] >= 1
    assert after["unique_clients"] == before["unique_clients"] == 1
    assert after["active_sessions"] == 0


async def test_report_export_strips_control_characters(world: World) -> None:
    from openpyxl import load_workbook

    world.client.name = "Shop\x07 One\x1f\x0b"
    await world.session.commit()
    await place(world, [("pb", 0, 1)])
    r = await world.admin.get("/api/admin/reports/sales.xlsx", params={"group_by": "client"})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))][1] == "Shop One"


@pytest.mark.parametrize(
    "params",
    [{"date_to": "9999-12-31"}, {"date_from": "1900-01-01", "date_to": "1900-01-02"}],
)
async def test_report_rejects_absurd_dates(world: World, params: dict[str, str]) -> None:
    assert (await world.admin.get("/api/admin/reports/sales", params=params)).status_code == 422


async def test_settings_audit_before_includes_telegram_without_secret(world: World) -> None:
    await _configure_telegram(world)
    await world.admin.put("/api/admin/settings", {"telegram": {"enabled": False}})
    last = (
        await world.session.scalars(
            select(AuditLog)
            .where(AuditLog.action == "settings.update")
            .order_by(AuditLog.id.desc())
        )
    ).first()
    assert last is not None and last.before is not None
    assert last.before["telegram"]["enabled"] is True
    assert last.before["telegram"]["chat_ids"] == ["42"]
    assert TOKEN not in str(last.before) and TOKEN not in str(last.after)
