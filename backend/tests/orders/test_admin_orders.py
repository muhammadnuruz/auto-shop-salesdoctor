"""Admin orders API: list/detail/export, edit (re-sends to SD), status, resend, delete, RBAC."""

import io
from collections.abc import Sequence
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select, update

from app.models import AuditLog, SdClient, ShopOrder
from app.services.jobs import RECONCILE_ORDER
from tests.api.conftest import Api, ApiFactory, make_settings
from tests.orders.conftest import WH, World

A = "/api/admin/orders"


async def sent(
    w: World, lines: Sequence[tuple[str, int, int | str]] | None = None
) -> dict[str, Any]:
    o = await w.draft(lines or [("pb", 0, 2)])
    assert (await w.submit(o)).status_code == 200
    got: dict[str, Any] = (await w.shop.get(f"/api/shop/orders/{o['public_id']}")).json()
    return got


async def admin_as(w: World, factory: ApiFactory, username: str, **kw: Any) -> Api:
    await w.seed.admin(username, **kw)
    api = await factory(make_settings(), sd=w.sd, jobs=w.jobs)
    assert (await api.admin_login(username)).status_code == 200
    return api


async def audit_actions(w: World) -> list[str]:
    await w.session.commit()
    return list((await w.session.scalars(select(AuditLog.action).order_by(AuditLog.id))).all())


# ------------------------------------------------------------------------------- list


async def test_list_summary_and_filters(world: World) -> None:
    a = await sent(world, [("pb", 0, 2)])  # 1000
    b = await sent(world, [("pa", 0, 1)])  # 1000
    await world.draft([("pb", 0, 1)])  # draft: 500
    body = (await world.admin.get(A)).json()
    assert body["total"] == 3 and body["summary"] == {"count": 3, "total_sum": 2500, "total_qty": 4}
    item = next(i for i in body["items"] if i["public_id"] == a["public_id"])
    assert item["client"]["name"] == "Тест Дилшод" and item["link"]["slug"] == "shop-link"
    assert item["agent"]["name"] == "AGENT11" and item["warehouse"]["id"] == WH
    only_sent = (await world.admin.get(A, params={"state": "sent"})).json()
    assert only_sent["total"] == 2 and only_sent["summary"]["total_sum"] == 2000
    assert (await world.admin.get(A, params={"state": ["draft", "failed"]})).json()["total"] == 1
    assert (await world.admin.get(A, params={"state": "zzz"})).status_code == 422
    by_code = (await world.admin.get(A, params={"q": b["code"]})).json()
    assert [i["public_id"] for i in by_code["items"]] == [b["public_id"]]
    by_client = (await world.admin.get(A, params={"q": "дилшод"})).json()
    assert by_client["total"] == 3
    assert (await world.admin.get(A, params={"client": world.client.id + 1})).json()["total"] == 0
    assert (await world.admin.get(A, params={"link": world.link.id})).json()["total"] == 3
    assert (await world.admin.get(A, params={"sd_status": 1})).json()["total"] == 2
    assert (await world.admin.get(A, params={"agent": "d0_3"})).json()["total"] == 0
    assert (await world.admin.get(A, params={"warehouse": WH})).json()["total"] == 3
    assert (await world.admin.get(A, params={"date_from": "2999-01-01"})).json()["total"] == 0
    page = (await world.admin.get(A, params={"limit": 1, "offset": 1})).json()
    assert len(page["items"]) == 1 and page["total"] == 3


async def test_detail_has_lines_events_and_refs(world: World) -> None:
    o = await sent(world, [("pa", 1, 1)])
    d = (await world.admin.get(f"{A}/{o['public_id']}")).json()
    assert (
        d["state"] == "sent"
        and d["client"]["sd_id"] == "w3_7448"
        and d["link"]["slug"] == "shop-link"
    )
    assert d["lines"][0]["quantity"] == 13 and d["lines"][0]["price"] == 1000
    names = [e["event"] for e in d["events"]]
    assert names[0] == "created" and names[-2:] == ["submitted", "sd_ok"]
    assert d["events"][0]["actor_type"] == "shop"
    assert (await world.admin.get(f"{A}/00000000-0000-0000-0000-000000000000")).status_code == 404


async def test_rbac_and_auth(world: World, app_factory: ApiFactory) -> None:
    o = await sent(world)
    anon = await app_factory(make_settings())
    assert (await anon.get(A)).status_code == 401
    assert (await world.shop.get(A)).status_code == 401  # shop sessions are not admin sessions
    viewer = await admin_as(world, app_factory, "view", role="viewer")
    assert (await viewer.get(A)).status_code == 200
    assert (await viewer.get(f"{A}/{o['public_id']}")).status_code == 200
    assert (await viewer.get(f"{A}/export.xlsx")).status_code == 200
    pid = o["public_id"]
    assert (await viewer.patch(f"{A}/{pid}", {"version": 1, "comment": "x"})).status_code == 403
    assert (await viewer.post(f"{A}/{pid}/status", {"status": 2})).status_code == 403
    assert (await viewer.post(f"{A}/{pid}/resend")).status_code == 403
    assert (await viewer.delete(f"{A}/{pid}")).status_code == 403
    writer = await admin_as(world, app_factory, "writer", role="viewer", extra=["orders.write"])
    assert (
        await writer.post(f"{A}/{pid}/status", {"status": 5})
    ).status_code == 403  # needs cancel
    assert (await writer.delete(f"{A}/{pid}")).status_code == 403
    assert (await writer.post(f"{A}/{pid}/status", {"status": 2})).status_code == 200
    assert world.sd_server.orders[o["code"]]["status"] == 2


async def test_csrf_required_for_admin_writes(world: World) -> None:
    o = await sent(world)
    r = await world.admin.post(f"{A}/{o['public_id']}/status", {"status": 2}, csrf=False)
    assert r.status_code == 403


# ------------------------------------------------------------------------------- edit


async def test_edit_sent_order_resends_same_code_and_updates_sd(world: World) -> None:
    o = await sent(world, [("pb", 0, 2)])
    assert world.sd_server.stock[WH]["pb"] == 3
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {
            "version": o["version"],
            "comment": "by phone",
            "lines": [{"product_ref": "pb", "pieces": "3"}, {"product_ref": "pc1", "pieces": "1"}],
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["comment"] == "by phone" and body["state"] == "sent"
    assert body["totals"]["qty"] == 4 and body["version"] == o["version"] + 1
    resent = world.sd_server.set_orders_sent()[-1]
    assert resent["code_1C"] == o["code"] and resent["comment"] == "by phone"
    assert [(p["product"]["SD_id"], p["quantity"]) for p in resent["orderProducts"]] == [
        ("pb", 3.0),
        ("pc1", 1.0),
    ]
    assert len(world.sd_server.orders) == 1  # updated, not duplicated
    assert world.sd_server.stock[WH]["pb"] == 2  # SD adjusted by the difference
    assert any(e["event"] == "admin_edit" for e in body["events"])
    assert "order.admin_edit" in await audit_actions(world)


async def test_admin_may_use_any_active_sd_reference(world: World) -> None:
    o = await sent(world)
    r = await world.admin.patch(
        f"{A}/{o['public_id']}", {"version": o["version"], "agent": "d0_3", "price_type": "d0_4"}
    )
    assert r.status_code == 200, r.text
    resent = world.sd_server.set_orders_sent()[-1]
    assert resent["agent"] == {"SD_id": "d0_3"} and resent["priceType"] == {"SD_id": "d0_4"}
    assert r.json()["lines"][0]["price"] == 450  # repriced for the new price type
    # but never an inactive / purchase reference
    version = r.json()["version"]
    for bad in ({"agent": "ghost"}, {"price_type": "d0_7"}, {"warehouse": "d0_9"}):
        r = await world.admin.patch(f"{A}/{o['public_id']}", {"version": version, **bad})
        assert r.status_code == 422, bad


async def test_admin_edit_credits_own_reservation(world: World) -> None:
    o = await sent(world, [("pb", 0, 5)])  # takes all 5; SD stock is now 0
    assert world.sd_server.stock[WH]["pb"] == 0
    ok = await world.admin.patch(
        f"{A}/{o['public_id']}", {"version": o["version"], "comment": "same qty"}
    )
    assert ok.status_code == 200  # 5 requested <= 0 available + 5 reserved by itself
    over = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": ok.json()["version"], "lines": [{"product_ref": "pb", "pieces": "6"}]},
    )
    assert over.status_code == 422
    assert over.json()["detail"]["errors"][0]["code"] == "over_stock"


async def test_admin_edit_sd_rejection_changes_nothing(world: World) -> None:
    o = await sent(world, [("pb", 0, 2)])
    world.sd_server.script = ["validation"]
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": o["version"], "lines": [{"product_ref": "pb", "pieces": "4"}]},
    )
    assert r.status_code == 422 and r.json()["detail"]["code"] == "sd_rejected"
    d = (await world.admin.get(f"{A}/{o['public_id']}")).json()
    assert d["version"] == o["version"] and d["lines"][0]["quantity"] == 2 and d["state"] == "sent"


async def test_admin_edit_unknown_outcome_keeps_edit_and_reconciles(world: World) -> None:
    o = await sent(world, [("pb", 0, 2)])
    world.sd_server.script = ["timeout"]
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": o["version"], "lines": [{"product_ref": "pb", "pieces": "4"}]},
    )
    assert r.status_code == 202 and r.json()["state"] == "outcome_unknown"
    assert r.json()["lines"][0]["quantity"] == 4
    assert any(name == RECONCILE_ORDER for name, _, _ in world.jobs.enqueued)


async def test_admin_edit_rules(world: World) -> None:
    o = await sent(world)
    stale = await world.admin.patch(f"{A}/{o['public_id']}", {"version": 1, "comment": "x"})
    assert stale.status_code == 409 and stale.json()["detail"]["code"] == "version_conflict"
    world.sd_server.orders[o["code"]]["status"] = 2
    from app.services import order_jobs

    await order_jobs.sync_order_statuses(world.sync)  # SD says shipped now
    r = await world.admin.patch(f"{A}/{o['public_id']}", {"version": o["version"], "comment": "x"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "order_not_editable"


async def test_admin_edit_failed_order_is_local_only(world: World) -> None:
    o = await world.draft([("pa", 0, 1)])
    world.sd_server.script = ["validation"]
    await world.submit(o)
    before = len(world.sd_server.calls)
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": 99, "lines": [{"product_ref": "pb", "pieces": "1"}]},
    )
    assert r.status_code == 409
    d = (await world.admin.get(f"{A}/{o['public_id']}")).json()
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": d["version"], "lines": [{"product_ref": "pb", "pieces": "1"}]},
    )
    assert r.status_code == 200 and r.json()["state"] == "failed"
    assert len(world.sd_server.calls) == before  # nothing sent until a resend


async def test_admin_edit_applies_same_line_validation(world: World) -> None:
    o = await sent(world)
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": o["version"], "lines": [{"product_ref": "pb", "blocks": 1, "pieces": "0"}]},
    )
    assert r.status_code == 422
    r = await world.admin.patch(
        f"{A}/{o['public_id']}",
        {"version": o["version"], "lines": [{"product_ref": "pnoprice", "pieces": "1"}]},
    )
    assert r.status_code == 422


# ------------------------------------------------------------------------------ status


async def test_status_change_sends_set_status(world: World) -> None:
    o = await sent(world)
    r = await world.admin.post(
        f"{A}/{o['public_id']}/status", {"status": 2, "shipment_date": "2026-10-20"}
    )
    assert r.status_code == 200, r.text
    assert r.json()["sd_status"] == 2 and r.json()["state"] == "sent"
    (call,) = world.sd_server.calls_of("setStatus")
    (change,) = call["data"]["order"]
    assert change["code_1C"] == o["code"] and change["status"] == 2
    assert change["dateShipment"] == "2026-10-20"
    assert "order.status" in await audit_actions(world)
    assert [e["event"] for e in r.json()["events"]][-1] == "status_changed"


async def test_status_5_cancels_and_warns_after_shipment(world: World) -> None:
    o = await sent(world)
    await world.admin.post(f"{A}/{o['public_id']}/status", {"status": 2})
    r = await world.admin.post(f"{A}/{o['public_id']}/status", {"status": 5})
    assert r.status_code == 200 and r.json()["state"] == "cancelled"
    assert r.json()["warnings"] == ["cancel_after_shipment_creates_return"]
    back = await world.admin.post(f"{A}/{o['public_id']}/status", {"status": 1})
    assert back.status_code == 409  # cancelled orders are not "sent" any more


async def test_status_validation(world: World) -> None:
    o = await sent(world)
    for bad in (0, 6, "x"):
        assert (
            await world.admin.post(f"{A}/{o['public_id']}/status", {"status": bad})
        ).status_code == 422
    draft = await world.draft()
    r = await world.admin.post(f"{A}/{draft['public_id']}/status", {"status": 2})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "invalid_state"


# ------------------------------------------------------------------------ resend / delete


async def test_resend_failed_order_runs_full_submit(world: World) -> None:
    o = await world.draft([("pb", 0, 1)])
    world.sd_server.script = ["validation"]
    await world.submit(o)
    r = await world.admin.post(f"{A}/{o['public_id']}/resend")
    assert r.status_code == 200, r.text
    assert r.json()["state"] == "sent"
    assert len(world.sd_server.orders) == 1
    again = await world.admin.post(f"{A}/{o['public_id']}/resend")
    assert again.status_code == 409  # nothing to resend


async def test_resend_unknown_order_reconciles(world: World) -> None:
    o = await world.draft([("pb", 0, 1)])
    world.sd_server.script = ["timeout_applied"]
    await world.submit(o)
    r = await world.admin.post(f"{A}/{o['public_id']}/resend")
    assert r.status_code == 200 and r.json()["reconcile"] == "sent" and r.json()["state"] == "sent"
    assert len(world.sd_server.calls_of("setOrder")) == 1


async def test_resend_unknown_sd_down_is_pending(world: World) -> None:
    o = await world.draft([("pb", 0, 1)])
    world.sd_server.script = ["timeout"]
    await world.submit(o)
    world.sd_server.orders_down = True
    r = await world.admin.post(f"{A}/{o['public_id']}/resend")
    assert r.status_code == 202 and r.json()["reconcile"] == "pending"


async def test_delete_sent_order_cancels_then_deletes_in_sd(world: World) -> None:
    o = await sent(world, [("pb", 0, 2)])
    r = await world.admin.delete(f"{A}/{o['public_id']}")
    assert r.status_code == 200, r.text
    assert r.json()["state"] == "deleted"
    assert [s["status"] for s in world.sd_server.set_orders_sent()] == [1, 5]  # stock restored
    assert world.sd_server.deleted == [o["code"]]
    assert world.sd_server.stock[WH]["pb"] == 5
    assert "order.delete" in await audit_actions(world)
    assert (await world.shop.get("/api/shop/orders")).json()["total"] == 0  # hidden from the shop
    assert (await world.admin.delete(f"{A}/{o['public_id']}")).status_code == 409


async def test_delete_draft_does_not_touch_sd(world: World) -> None:
    o = await world.draft()
    r = await world.admin.delete(f"{A}/{o['public_id']}")
    assert r.status_code == 200 and r.json()["state"] == "deleted"
    assert world.sd_server.calls == []


async def test_delete_sd_failure_keeps_order(world: World) -> None:
    o = await sent(world)
    world.sd_server.script = ["connect_error", "connect_error"]
    r = await world.admin.delete(f"{A}/{o['public_id']}")
    assert r.status_code == 502
    assert (await world.admin.get(f"{A}/{o['public_id']}")).json()["state"] == "sent"
    assert world.sd_server.deleted == []


# ----------------------------------------------------------------------------- export


async def test_export_xlsx(world: World) -> None:
    await sent(world)
    await world.session.execute(update(SdClient).values(name='=HYPERLINK("x")'))
    await world.session.commit()
    r = await world.admin.get(f"{A}/export.xlsx", params={"state": "sent"})
    assert r.status_code == 200 and r.headers["X-Rows"] == "1"
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Code" and rows[1][0].startswith("WEB-") and rows[1][2] == "sent"
    assert rows[1][4] == '\'=HYPERLINK("x")'  # formula injection neutralised
    assert rows[1][12] == 1000
    assert "order.export" in await audit_actions(world)
    empty = await world.admin.get(f"{A}/export.xlsx", params={"state": "failed"})
    assert empty.headers["X-Rows"] == "0"


async def test_order_row_counts_in_db(world: World) -> None:
    await sent(world)
    await world.session.commit()
    assert len((await world.session.scalars(select(ShopOrder))).all()) == 1
